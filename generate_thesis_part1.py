# -*- coding: utf-8 -*-
"""
宠物寄养登记与健康档案管理系统 - 论文初稿生成脚本
基于实际开发的系统完整改写
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# ========== 全局样式设置 ==========
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(12)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = '黑体'
    hs._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    hs.font.color.rgb = RGBColor(0, 0, 0)

doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(14)
doc.styles['Heading 3'].font.size = Pt(13)

def add_para(text, bold=False, align=None, size=None, first_indent=True):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    return p

def add_heading(text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    return h

def add_table(headers, rows, caption=None):
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.bold = True
        run.font.size = Pt(10.5)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10.5)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10.5)
    doc.add_paragraph()
    return table

# ========== 封面 ==========
for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('毕 业 设 计')
run.font.size = Pt(26)
run.bold = True
run.font.name = '黑体'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

doc.add_paragraph()
doc.add_paragraph()

cover_items = [
    ('中文题目', '宠物寄养登记与健康档案管理系统的设计与实现'),
    ('英文题目', 'Design and Implementation of a Pet Boarding Registration and Health Record Management System'),
    ('学    院', '信息工程学院'),
    ('年级专业', '2023级信息管理与信息系统'),
    ('姓    名', '陈俊斌'),
    ('学    号', '232022226'),
    ('指导教师', '朱怡贺'),
    ('职    称', '助教'),
]
for label, value in cover_items:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'{label}：{value}')
    run.font.size = Pt(14)

doc.add_paragraph()
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('闽南理工学院教务处制')
run.font.size = Pt(14)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('2027年5月10日')
run.font.size = Pt(14)

doc.add_page_break()

# ========== 诚信声明 ==========
add_heading('毕业设计诚信声明书', 1)
add_para('本人郑重声明：在毕业设计工作中严格遵守学校有关规定，恪守学术规范；我所提交的毕业设计是本人在朱怡贺指导教师的指导下独立研究、撰写的成果，设计中所引用他人的文字、研究成果，均已在设计中加以说明；在本人的毕业设计中未剽窃、抄袭他人的学术观点、思想和成果，未篡改实验数据。')
add_para('本设计和资料若有不实之处，本人愿承担一切相关责任。')
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run('学生签名：陈俊斌')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run('2027年5月10日')

doc.add_page_break()

# ========== 中文摘要 ==========
add_heading('宠物寄养登记与健康档案管理系统的设计与实现', 1)
add_para('【摘 要】随着我国宠物经济的快速发展和居民生活水平的不断提高，宠物寄养服务需求日益增长，传统的人工登记管理方式已难以满足宠物寄养门店的运营需求。宠物寄养登记与健康档案管理系统是面向宠物寄养门店的信息化管理平台，旨在实现寄养业务全流程的数字化管理，提高门店运营效率和服务质量。系统采用前后端分离架构，前端基于React、TypeScript、Vite和Tailwind CSS构建，使用shadcn/ui组件库实现统一的界面风格；后端采用Spring Boot 3.2框架，结合MyBatis-Plus持久层框架和MySQL 8.0数据库，使用JWT实现无状态身份认证。系统主要包含客户信息管理、宠物档案管理、寄养订单管理、健康档案管理、寄养服务记录、统计分析和系统管理七大功能模块，支持管理员和工作人员两种角色权限。系统实现了寄养订单从创建、确认入店、寄养中到确认离店的全生命周期管理，宠物健康档案的持续追踪，以及多维度的经营数据统计分析。经过功能测试和E2E自动化测试验证，系统运行稳定，功能完整，能够有效提升宠物寄养门店的管理效率和服务水平。', first_indent=False)
add_para('【关键词】宠物寄养，健康档案，前后端分离，Spring Boot，React', first_indent=False)

doc.add_page_break()

# ========== 英文摘要 ==========
add_heading('Design and Implementation of a Pet Boarding Registration and Health Record Management System', 1)
add_para('【Abstract】With the rapid development of China\'s pet economy and the continuous improvement of residents\' living standards, the demand for pet boarding services is growing day by day. Traditional manual registration management methods can no longer meet the operational needs of pet boarding stores. The Pet Boarding Registration and Health Record Management System is an information management platform for pet boarding stores, aiming to realize digital management of the entire boarding business process and improve store operational efficiency and service quality. The system adopts a front-end and back-end separation architecture. The front-end is built based on React, TypeScript, Vite, and Tailwind CSS, using the shadcn/ui component library to achieve a unified interface style. The back-end uses the Spring Boot 3.2 framework, combined with the MyBatis-Plus persistence layer framework and MySQL 8.0 database, and uses JWT to implement stateless identity authentication. The system mainly includes seven functional modules: customer information management, pet profile management, boarding order management, health record management, boarding service records, statistical analysis, and system management, supporting two role permissions: administrator and staff. The system realizes full lifecycle management of boarding orders from creation, check-in confirmation, during boarding, to check-out confirmation, continuous tracking of pet health records, and multi-dimensional business data statistical analysis. Verified through functional testing and E2E automated testing, the system runs stably with complete functions, and can effectively improve the management efficiency and service level of pet boarding stores.', first_indent=False)
add_para('【Keywords】Pet boarding, Health record, Front-end and back-end separation, Spring Boot, React', first_indent=False)

doc.add_page_break()

# ========== 目录 ==========
add_heading('目录', 1)
toc_items = [
    '1.绪论',
    '1.1系统开发的背景、目的和意义',
    '1.1.1系统开发的背景',
    '1.1.2系统开发的目的和意义',
    '1.2国内外研究现状',
    '2.系统分析',
    '2.1可行性分析',
    '2.1.1经济可行性',
    '2.1.2技术可行性',
    '2.1.3操作可行性',
    '2.2需求分析',
    '2.3业务流程分析',
    '2.4数据流程分析',
    '3.系统设计',
    '3.1系统功能设计',
    '3.2数据库设计',
    '3.2.1数据库概念结构设计',
    '3.2.2数据库物理结构设计',
    '4.系统实现',
    '4.1系统主界面',
    '4.2客户信息管理模块',
    '4.3宠物档案管理模块',
    '4.4寄养订单管理模块',
    '4.5健康档案管理模块',
    '4.6寄养服务记录模块',
    '4.7统计分析模块',
    '4.8系统管理模块',
    '5.系统测试',
    '5.1系统测试目的和意义',
    '5.2系统测试方法',
    '5.3系统功能测试用例',
    '6.总结',
    '参考文献',
    '致谢',
]
for item in toc_items:
    p = doc.add_paragraph()
    run = p.add_run(item)
    run.font.size = Pt(12)

doc.add_page_break()

# ========== 第1章 绪论 ==========
add_heading('1.绪论', 1)

add_heading('1.1系统开发的背景、目的和意义', 2)

add_heading('1.1.1系统开发的背景', 3)
add_para('近年来，随着我国经济的持续发展和居民可支配收入的稳步增长，宠物饲养已从少数人的爱好转变为普遍的生活方式。据《2025年中国宠物行业白皮书》数据显示，我国城镇宠物犬猫数量已超过1.2亿只，宠物消费市场规模突破3000亿元，其中宠物寄养服务作为宠物消费的重要组成部分，呈现出快速增长的态势。')
add_para('在宠物寄养行业蓬勃发展的同时，传统的门店管理方式也暴露出诸多问题。目前，大多数中小型宠物寄养门店仍采用纸质登记本或简单的Excel表格来管理客户信息、宠物档案和寄养订单，这种方式存在以下不足：一是信息检索效率低下，当客户数量增多时，查找特定宠物的寄养记录和健康档案需要耗费大量时间；二是数据容易丢失或损坏，纸质登记本易磨损，Excel文件存在误删和版本混乱的风险；三是寄养状态跟踪不及时，宠物入店、离店、健康检查等关键节点缺乏系统化的记录和提醒；四是经营数据分析困难，无法直观了解门店的订单量、收入趋势、宠物品种分布等关键经营指标。')
add_para('与此同时，互联网技术、云计算和移动终端的快速发展为宠物寄养行业的信息化转型提供了技术支撑。Spring Boot、React等成熟的前后端开发框架，MySQL等关系型数据库，以及JWT等安全认证技术，使得开发一套功能完善、运行稳定的宠物寄养管理系统变得切实可行。基于上述背景，开发一套宠物寄养登记与健康档案管理系统，对于提升宠物寄养门店的管理效率和服务质量具有重要的现实意义。')

add_heading('1.1.2系统开发的目的和意义', 3)
add_para('本系统的开发目的在于为宠物寄养门店提供一套功能完善、操作便捷的信息化管理平台，实现寄养业务全流程的数字化管理。具体而言，系统的开发目的和意义体现在以下几个方面：')
add_para('第一，提高门店运营效率。通过系统化的客户信息管理、宠物档案管理和寄养订单管理，门店工作人员可以快速完成客户登记、宠物建档、订单创建等操作，大幅减少人工录入和信息检索的时间成本。系统支持寄养订单从创建、确认入店、寄养中到确认离店的全生命周期管理，每个状态变更都有明确的操作记录和时间戳，确保业务流程的规范化和可追溯性。')
add_para('第二，提升宠物健康管理水平。系统为每只宠物建立独立的健康档案，记录宠物的疫苗接种、体检结果、用药记录、异常情况等健康信息。在寄养期间，工作人员可以随时登记宠物的健康状况，系统会自动关联到对应宠物的健康档案中，实现宠物健康数据的持续积累和追踪。这不仅有助于门店及时发现和处理宠物的健康问题，也为宠物主人提供了透明、可信的健康记录。')
add_para('第三，辅助经营决策。系统的统计分析模块提供多维度的经营数据可视化，包括寄养订单趋势、宠物品种分布、客户消费统计、月度收入分析等。门店管理者可以通过直观的图表了解门店的经营状况，发现业务增长机会和潜在问题，从而做出更加科学的经营决策。')
add_para('第四，改善客户服务体验。系统化的管理使得客户信息和宠物档案得以完整保存，当老客户再次到店时，工作人员可以快速调取历史记录，提供更加个性化和贴心的服务。同时，规范的寄养流程和透明的健康记录也有助于增强客户对门店的信任感和满意度。')

add_heading('1.2国内外研究现状', 2)
add_para('在国外，宠物寄养行业的信息化程度相对较高，欧美等发达国家的宠物寄养门店普遍采用专业化的管理软件进行日常运营。例如，美国的PetExec、德国的PetAdmin等宠物寄养管理软件，提供了在线预约、客户管理、宠物档案、寄养日程、账单支付等功能，部分软件还支持移动端应用和客户自助门户。这些系统通常采用SaaS（软件即服务）模式，门店按月或按年订阅使用，降低了门店的信息化门槛。在技术架构上，国外的宠物管理系统普遍采用云原生架构，支持多租户部署，具备良好的可扩展性和数据安全性。')
add_para('在国内，随着宠物经济的快速发展，宠物寄养管理系统也逐渐受到关注。目前市场上的宠物管理软件主要分为两类：一类是面向大型连锁宠物医院和寄养中心的综合管理系统，如瑞鹏宠物医院的内部管理系统、新瑞鹏集团的数字化管理平台等，这类系统功能全面但部署成本较高，主要服务于大型企业客户；另一类是面向中小型宠物门店的轻量化管理工具，如宠物管家、宠老板等SaaS平台，提供基础的客户管理、寄养登记和收银功能，但在健康档案管理和数据分析方面相对薄弱。')
add_para('从学术研究角度来看，国内学者在宠物管理系统方面的研究主要集中在系统设计与实现层面。相关研究普遍采用Java Web技术栈，结合SSM或Spring Boot框架开发宠物寄养或宠物医院管理系统，实现了客户管理、宠物信息管理、寄养预约等基础功能。然而，现有研究在以下方面仍存在不足：一是多数系统侧重于寄养预约和订单管理，对宠物健康档案的持续追踪和深度管理关注不够；二是系统的角色权限设计较为简单，缺乏对门店不同岗位工作人员的精细化权限控制；三是经营数据分析功能薄弱，多数系统仅提供基础的数据列表，缺乏多维度的可视化统计分析；四是前端界面设计较为传统，用户体验有待提升。')
add_para('基于对国内外研究现状的分析，本系统在设计和实现上进行了针对性的改进：一是建立完善的宠物健康档案管理模块，实现宠物健康数据的全生命周期追踪；二是设计管理员和工作人员两种角色，实现精细化的权限控制；三是开发多维度的统计分析模块，提供直观的数据可视化；四是采用现代化的前端技术栈和组件库，打造简洁美观、操作便捷的用户界面。')

doc.add_page_break()

# ========== 第2章 系统分析 ==========
add_heading('2.系统分析', 1)

add_heading('2.1可行性分析', 2)

add_heading('2.1.1经济可行性', 3)
add_para('本系统采用B/S架构设计，用户通过浏览器即可访问系统，无需安装专用客户端软件，降低了系统的部署和维护成本。在开发阶段，系统所使用的技术栈均为开源免费技术：前端采用React、TypeScript、Vite和Tailwind CSS等开源框架和工具，后端采用Spring Boot、MyBatis-Plus等开源框架，数据库采用MySQL社区版，这些技术均无需支付授权费用。在部署阶段，系统可以部署在普通的云服务器上，对于中小型宠物寄养门店而言，一台基础配置的云服务器即可满足运行需求，月度运维成本较低。此外，系统的投入使用将显著提高门店的运营效率，减少人工管理成本，从长远来看具有良好的经济效益。因此，从经济角度分析，开发本系统是完全可行的。')

add_heading('2.1.2技术可行性', 3)
add_para('本系统的开发采用成熟稳定的前后端分离技术栈。前端方面，系统使用React 19作为UI框架，React具有组件化开发、虚拟DOM、单向数据流等特性，能够高效构建复杂的用户界面；使用TypeScript提供静态类型检查，增强代码的可维护性和可靠性；使用Vite作为构建工具，提供极速的开发服务器启动和热模块替换；使用Tailwind CSS实现原子化CSS样式开发，配合shadcn/ui组件库，能够快速构建风格统一、美观现代的用户界面。后端方面，系统采用Spring Boot 3.2框架，Spring Boot通过自动配置和起步依赖大幅简化了Spring应用的搭建和开发过程；使用MyBatis-Plus作为持久层框架，在MyBatis的基础上提供了强大的CRUD操作封装和条件构造器，减少了大量重复的SQL编写工作；使用MySQL 8.0作为关系型数据库，MySQL具有性能优异、稳定可靠、社区活跃等优点，能够满足系统的数据存储需求；使用JWT（JSON Web Token）实现无状态身份认证，便于系统的水平扩展。上述技术均为当前业界主流技术，文档丰富，社区活跃，开发人员能够快速上手。因此，从技术角度分析，开发本系统是完全可行的。')

add_heading('2.1.3操作可行性', 3)
add_para('本系统的用户界面设计遵循简洁直观的原则，采用侧边栏导航加内容区域的经典后台管理系统布局，用户可以通过左侧导航菜单快速切换到各个功能模块。系统的操作流程设计贴合宠物寄养门店的实际业务场景，工作人员经过简单的培训即可上手使用。对于常见的操作，如新增客户、创建订单、登记健康记录等，系统提供了清晰的表单界面和明确的操作按钮，降低了用户的学习成本。此外，系统采用响应式设计，在不同分辨率的屏幕上都能正常显示和操作。因此，从操作角度分析，本系统具有良好的易用性，开发和推广是完全可行的。')

add_heading('2.2需求分析', 2)
add_para('宠物寄养登记与健康档案管理系统的目标用户是宠物寄养门店的管理者和工作人员。系统旨在为门店提供一个高效、便捷的信息化管理平台，覆盖客户管理、宠物档案、寄养订单、健康记录、服务记录、统计分析和系统管理等核心业务场景。通过对宠物寄养门店的实际业务流程进行调研和分析，将系统的功能需求划分为以下模块：')
add_para('（1）系统首页/工作台：作为用户登录系统后的首页面，展示门店的核心经营数据和待办事项。包括寄养中宠物数量、待入店订单数量、本月订单数量、本月寄养金额等统计卡片，今日待办订单列表，在店宠物状态概览，以及常用业务操作的快捷入口。')
add_para('（2）客户信息管理：管理宠物主人的联系信息。支持客户信息的新增、编辑、删除和查询，记录客户的姓名、手机号码、地址、备注等信息。系统对客户的手机号码进行唯一性校验，避免重复登记。删除客户时，如果该客户名下有关联的宠物或订单，系统将阻止删除操作，确保数据的完整性。')
add_para('（3）宠物档案管理：管理宠物的基本信息和档案资料。支持宠物信息的新增、编辑、删除和查询，记录宠物的名称、品种、性别、年龄、体重、毛色、绝育情况、疫苗情况、过敏史、特殊需求等信息。每只宠物关联到对应的客户，一个客户可以拥有多只宠物。')
add_para('（4）寄养订单管理：管理宠物寄养的订单全流程。支持寄养订单的创建、编辑、查询和状态管理。订单记录包含订单编号、客户信息、宠物信息、寄养开始时间、结束时间、寄养类型、服务项目、费用明细、订单状态等信息。订单状态包括待入店、寄养中、已完成、已取消四种，系统支持订单状态的流转操作：确认入店（待入店→寄养中）、确认离店（寄养中→已完成）、取消订单（待入店→已取消，仅管理员可操作）。系统自动生成唯一的订单编号，便于订单的识别和追踪。')
add_para('（5）健康档案管理：管理宠物的健康记录。支持健康记录的新增、编辑、查询和删除，记录宠物的体温、体重、精神状态、食欲情况、排便情况、异常描述、处理措施、记录时间等信息。每条健康记录关联到对应的宠物和寄养订单，系统自动记录操作人和操作时间。为保证数据的准确性，系统限制当天的健康记录可以编辑，非当天的记录仅支持查看。')
add_para('（6）寄养服务记录：管理寄养期间提供的附加服务记录。支持服务记录的新增、编辑、查询和删除，记录服务项目、服务内容、服务时间、收费金额、操作人等信息。服务记录关联到对应的寄养订单，便于订单费用的核算和统计。')
add_para('（7）统计分析：提供多维度的经营数据统计和可视化分析。包括寄养订单趋势图（按月统计订单数量和寄养收入）、宠物品种分布环形图、统计卡片（订单总数、已完成订单、已取消订单、总营收）。支持按时间范围（近7天、近30天、本月、本季度、本年、自定义）筛选数据，用户可以灵活选择统计周期。系统支持订单数据的Excel导出，方便门店进行数据备份和进一步分析。')
add_para('（8）系统管理：管理系统的用户账号和权限。支持用户的新增、编辑、删除、启用/禁用操作，管理用户的账号、密码、真实姓名、角色、备注等信息。系统设计了管理员和工作人员两种角色：管理员拥有系统的全部权限，包括用户管理、订单取消等高级操作；工作人员拥有业务操作权限，但不能访问系统管理模块，也不能取消订单。系统支持用户修改个人密码，记录用户的登录时间和操作日志。')
add_para('除上述功能需求外，系统还需满足以下非功能性需求：一是安全性需求，系统采用JWT身份认证机制，用户登录后获取令牌，后续请求携带令牌进行身份验证，未登录用户无法访问系统资源；系统对用户密码进行加密存储，防止密码泄露；系统实现基于角色的访问控制，不同角色的用户只能访问其权限范围内的功能。二是性能需求，系统应具备良好的响应速度，常规页面加载时间不超过3秒，数据查询操作响应及时。三是可靠性需求，系统应稳定运行，具备异常处理能力，在出现错误时给出友好的提示信息，避免系统崩溃。四是可维护性需求，系统采用模块化设计，代码结构清晰，便于后续的功能扩展和维护。')

add_heading('2.3业务流程分析', 2)
add_para('宠物寄养门店的核心业务流程围绕寄养订单的全生命周期展开。客户携带宠物到店后，工作人员首先查询客户信息，如果是新客户则先进行客户登记，然后为宠物建立档案。客户和宠物信息确认无误后，工作人员创建寄养订单，填写寄养时间、服务项目、费用明细等信息，订单创建后状态为"待入店"。')
add_para('当宠物正式入店时，工作人员执行"确认入店"操作，订单状态变更为"寄养中"。在寄养期间，工作人员需要定期为宠物进行健康检查并登记健康记录，根据客户需求提供附加服务并登记服务记录。如果发现宠物健康异常，工作人员应及时记录并通知客户，必要时联系宠物医院进行诊治。')
add_para('当寄养结束客户接回宠物时，工作人员执行"确认离店"操作，订单状态变更为"已完成"，同时进行费用结算。如果客户在宠物入店前取消寄养计划，管理员可以执行"取消订单"操作，订单状态变更为"已取消"。')
add_para('在整个业务流程中，系统对每个关键节点的操作都进行记录，包括操作人、操作时间和操作内容，确保业务流程的可追溯性。同时，系统通过统计分析模块对订单数据进行汇总分析，为门店管理者提供经营决策支持。')

add_heading('2.4数据流程分析', 2)
add_para('在本系统中，主要涉及管理员和工作人员两种外部主体。工作人员负责日常业务操作，包括客户登记、宠物建档、订单创建、健康记录登记、服务记录登记等；管理员除了具备工作人员的全部业务权限外，还负责系统用户管理、订单取消、权限分配等高级管理操作。')
add_para('从数据流向来看，工作人员在客户端录入客户信息、宠物信息、订单信息、健康记录和服务记录，这些数据通过网络传输到服务器端，服务器端进行业务逻辑处理后将数据持久化到MySQL数据库中。当工作人员需要查询数据时，查询请求发送到服务器端，服务器端从数据库中检索数据并返回给客户端展示。管理员可以查看所有业务数据和系统操作日志，并进行用户管理和系统配置操作。')
add_para('系统的数据流程可以分为以下几个主要子流程：一是客户与宠物管理子流程，工作人员录入客户信息和宠物档案，系统进行唯一性校验和数据关联，将数据保存到数据库；二是寄养订单管理子流程，工作人员创建寄养订单，系统自动生成订单编号，记录订单状态变更历史，关联客户和宠物信息；三是健康与服务记录子流程，工作人员在寄养期间登记宠物健康记录和服务记录，系统自动关联到对应订单和宠物，记录操作人和时间戳；四是统计分析子流程，系统从数据库中聚合订单数据，生成趋势图、分布图和统计指标，支持按时间范围筛选；五是系统管理子流程，管理员管理用户账号和角色权限，系统记录用户登录信息和操作日志。')

doc.add_page_break()

# 保存第一部分
doc.save(r'C:\Users\chen\Desktop\pet-foster\宠物寄养登记与健康档案管理系统_论文初稿.docx')
print('第一部分生成完成')
