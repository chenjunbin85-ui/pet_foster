# -*- coding: utf-8 -*-
"""论文第二部分：系统设计 + 系统实现"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document(r'C:\Users\chen\Desktop\pet-foster\宠物寄养登记与健康档案管理系统_论文初稿.docx')

def add_para(text, bold=False, align=None, size=None, first_indent=True):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    return p

def add_heading(text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    return h

def add_code(code_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    return p

def add_table(headers, rows, caption=None):
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.bold = True
        run.font.size = Pt(10.5)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10.5)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10.5)
    doc.add_paragraph()
    return table

# ========== 第3章 系统设计 ==========
add_heading('3.系统设计', 1)

add_heading('3.1系统功能设计', 2)
add_para('根据需求分析的结果，宠物寄养登记与健康档案管理系统采用基于角色的功能设计，系统用户分为管理员和工作人员两种角色。管理员拥有系统的全部功能权限，包括系统首页、客户信息管理、宠物档案管理、寄养订单管理、健康档案管理、寄养服务记录、统计分析和系统管理八大模块。工作人员拥有业务操作权限，可以访问系统首页、客户信息管理、宠物档案管理、寄养订单管理、健康档案管理、寄养服务记录和统计分析模块，但不能访问系统管理模块，也不能执行取消订单等高级操作。')
add_para('系统的整体功能结构如下：系统首页作为工作台，展示门店经营数据概览和待办事项；客户信息管理模块负责客户资料的维护；宠物档案管理模块负责宠物基本信息的管理；寄养订单管理模块是系统的核心模块，负责寄养订单的全生命周期管理；健康档案管理模块负责宠物健康记录的登记和查询；寄养服务记录模块负责附加服务的记录；统计分析模块提供多维度的经营数据分析；系统管理模块负责用户账号和权限的管理。各模块之间通过客户ID、宠物ID、订单ID等关键字段进行数据关联，形成完整的业务数据链条。')

add_heading('3.2数据库设计', 2)

add_heading('3.2.1数据库概念结构设计', 3)
add_para('根据系统的功能需求，设计了以下八个核心实体：')
add_para('（1）用户实体（sys_user）：用于存储系统用户的账号信息，属性包括用户ID、账号、密码、真实姓名、角色、账号状态、手机号码、邮箱、最后登录时间、创建时间、更新时间、备注等。用户实体是系统的身份认证基础，每个用户对应一个角色（管理员或工作人员）。')
add_para('（2）客户实体（customer）：用于存储宠物主人的联系信息，属性包括客户ID、姓名、手机号码、地址、创建时间、更新时间、备注等。客户实体是宠物寄养业务的服务对象，一个客户可以拥有多只宠物。')
add_para('（3）宠物实体（pet）：用于存储宠物的基本档案信息，属性包括宠物ID、客户ID、宠物名称、品种、性别、出生日期、体重、毛色、是否绝育、疫苗情况、过敏史、特殊需求、创建时间、更新时间、备注等。宠物实体关联到客户实体，是寄养服务的直接对象。')
add_para('（4）寄养订单实体（foster_order）：用于存储寄养订单的详细信息，属性包括订单ID、订单编号、客户ID、宠物ID、寄养开始时间、寄养结束时间、实际入店时间、实际离店时间、寄养类型、服务项目、总费用、订单状态、创建人、创建时间、更新时间、备注等。寄养订单实体是系统的核心实体，关联客户和宠物，记录寄养业务的全流程信息。')
add_para('（5）健康记录实体（pet_health_record）：用于存储宠物的健康检查记录，属性包括记录ID、宠物ID、订单ID、体温、体重、精神状态、食欲情况、排便情况、异常描述、处理措施、记录人、记录时间、创建时间、更新时间等。健康记录实体关联到宠物和寄养订单，记录寄养期间宠物的健康状况。')
add_para('（6）服务记录实体（foster_service_record）：用于存储寄养期间提供的附加服务信息，属性包括记录ID、订单ID、服务项目、服务内容、收费金额、服务时间、操作人、创建时间、更新时间等。服务记录实体关联到寄养订单，记录寄养期间的附加服务和收费情况。')
add_para('（7）操作日志实体（sys_oper_log）：用于记录用户的关键操作信息，属性包括日志ID、用户ID、用户名、操作模块、操作类型、操作描述、请求方法、请求参数、IP地址、操作时间等。操作日志实体用于系统的安全审计和操作追溯。')
add_para('（8）字典实体（sys_dict）：用于存储系统的字典数据，属性包括字典ID、字典类型、字典标签、字典值、排序、状态、备注等。字典实体用于管理系统中的枚举数据，如宠物品种、寄养类型、订单状态等，提高系统的可维护性。')
add_para('在上述实体中，客户与宠物之间是一对多关系，一个客户可以拥有多只宠物；客户与寄养订单之间是一对多关系，一个客户可以有多个寄养订单；宠物与寄养订单之间是一对多关系，一只宠物可以有多个寄养订单；寄养订单与健康记录之间是一对多关系，一个订单可以有多条健康记录；寄养订单与服务记录之间是一对多关系，一个订单可以有多条服务记录；用户与操作日志之间是一对多关系，一个用户可以有多条操作日志。')

add_heading('3.2.2数据库物理结构设计', 3)
add_para('本系统选用MySQL 8.0作为数据库管理系统，数据库名为pet_foster。数据库中各表的物理结构设计如下：')

add_para('（1）用户信息表sys_user，用于存储系统用户的账号信息，如表3-1所示：', first_indent=False)
add_table(
    ['字段名称', '数据类型', '字段说明'],
    [
        ['user_id', 'bigint', '用户ID，主键，自动增长'],
        ['username', 'varchar(50)', '登录账号，非空，唯一'],
        ['password', 'varchar(100)', '密码（BCrypt加密），非空'],
        ['real_name', 'varchar(50)', '真实姓名'],
        ['role', 'varchar(20)', '角色（admin管理员/staff工作人员），非空'],
        ['phone', 'varchar(20)', '手机号码'],
        ['email', 'varchar(100)', '邮箱'],
        ['status', 'tinyint', '账号状态（0正常 1禁用），默认0'],
        ['last_login_time', 'datetime', '最后登录时间'],
        ['create_time', 'datetime', '创建时间'],
        ['update_time', 'datetime', '更新时间'],
        ['remark', 'varchar(500)', '备注'],
    ],
    '表3-1 用户信息表'
)

add_para('（2）客户信息表customer，用于存储宠物主人的联系信息，如表3-2所示：', first_indent=False)
add_table(
    ['字段名称', '数据类型', '字段说明'],
    [
        ['customer_id', 'bigint', '客户ID，主键，自动增长'],
        ['name', 'varchar(50)', '客户姓名，非空'],
        ['phone', 'varchar(20)', '手机号码，非空，唯一'],
        ['address', 'varchar(200)', '地址'],
        ['create_time', 'datetime', '创建时间'],
        ['update_time', 'datetime', '更新时间'],
        ['remark', 'varchar(500)', '备注'],
    ],
    '表3-2 客户信息表'
)

add_para('（3）宠物档案表pet，用于存储宠物的基本信息，如表3-3所示：', first_indent=False)
add_table(
    ['字段名称', '数据类型', '字段说明'],
    [
        ['pet_id', 'bigint', '宠物ID，主键，自动增长'],
        ['customer_id', 'bigint', '所属客户ID，外键'],
        ['name', 'varchar(50)', '宠物名称，非空'],
        ['species', 'varchar(20)', '物种（dog/cat/other）'],
        ['breed', 'varchar(50)', '品种'],
        ['gender', 'varchar(10)', '性别（male/female）'],
        ['birth_date', 'date', '出生日期'],
        ['weight', 'decimal(5,2)', '体重（kg）'],
        ['color', 'varchar(50)', '毛色'],
        ['is_neutered', 'tinyint', '是否绝育（0否 1是）'],
        ['vaccine_info', 'varchar(500)', '疫苗情况'],
        ['allergy', 'varchar(500)', '过敏史'],
        ['special_needs', 'varchar(500)', '特殊需求'],
        ['create_time', 'datetime', '创建时间'],
        ['update_time', 'datetime', '更新时间'],
        ['remark', 'varchar(500)', '备注'],
    ],
    '表3-3 宠物档案表'
)

add_para('（4）寄养订单表foster_order，用于存储寄养订单的详细信息，如表3-4所示：', first_indent=False)
add_table(
    ['字段名称', '数据类型', '字段说明'],
    [
        ['order_id', 'bigint', '订单ID，主键，自动增长'],
        ['order_no', 'varchar(30)', '订单编号，非空，唯一'],
        ['customer_id', 'bigint', '客户ID，外键，非空'],
        ['pet_id', 'bigint', '宠物ID，外键，非空'],
        ['start_date', 'datetime', '寄养开始时间，非空'],
        ['end_date', 'datetime', '寄养结束时间，非空'],
        ['check_in_time', 'datetime', '实际入店时间'],
        ['check_out_time', 'datetime', '实际离店时间'],
        ['boarding_type', 'varchar(20)', '寄养类型'],
        ['services', 'varchar(500)', '服务项目'],
        ['total_fee', 'decimal(10,2)', '总费用'],
        ['status', 'varchar(20)', '订单状态（pending待入店/boarding寄养中/completed已完成/cancelled已取消），非空'],
        ['create_by', 'varchar(50)', '创建人'],
        ['create_time', 'datetime', '创建时间'],
        ['update_time', 'datetime', '更新时间'],
        ['remark', 'varchar(500)', '备注'],
    ],
    '表3-4 寄养订单表'
)

add_para('（5）健康记录表pet_health_record，用于存储宠物的健康检查记录，如表3-5所示：', first_indent=False)
add_table(
    ['字段名称', '数据类型', '字段说明'],
    [
        ['record_id', 'bigint', '记录ID，主键，自动增长'],
        ['pet_id', 'bigint', '宠物ID，外键，非空'],
        ['order_id', 'bigint', '关联订单ID，外键'],
        ['temperature', 'decimal(4,1)', '体温（℃）'],
        ['weight', 'decimal(5,2)', '体重（kg）'],
        ['spirit', 'varchar(20)', '精神状态'],
        ['appetite', 'varchar(20)', '食欲情况'],
        ['defecation', 'varchar(20)', '排便情况'],
        ['abnormal_desc', 'varchar(500)', '异常描述'],
        ['treatment', 'varchar(500)', '处理措施'],
        ['record_by', 'varchar(50)', '记录人'],
        ['record_time', 'datetime', '记录时间，非空'],
        ['create_time', 'datetime', '创建时间'],
        ['update_time', 'datetime', '更新时间'],
    ],
    '表3-5 健康记录表'
)

add_para('（6）服务记录表foster_service_record，用于存储寄养期间的附加服务记录，如表3-6所示：', first_indent=False)
add_table(
    ['字段名称', '数据类型', '字段说明'],
    [
        ['record_id', 'bigint', '记录ID，主键，自动增长'],
        ['order_id', 'bigint', '关联订单ID，外键，非空'],
        ['service_name', 'varchar(100)', '服务项目，非空'],
        ['description', 'varchar(500)', '服务内容'],
        ['fee', 'decimal(10,2)', '收费金额'],
        ['service_time', 'datetime', '服务时间'],
        ['operator', 'varchar(50)', '操作人'],
        ['create_time', 'datetime', '创建时间'],
        ['update_time', 'datetime', '更新时间'],
    ],
    '表3-6 服务记录表'
)

add_para('（7）操作日志表sys_oper_log，用于记录用户的关键操作，如表3-7所示：', first_indent=False)
add_table(
    ['字段名称', '数据类型', '字段说明'],
    [
        ['log_id', 'bigint', '日志ID，主键，自动增长'],
        ['user_id', 'bigint', '用户ID'],
        ['username', 'varchar(50)', '用户名'],
        ['module', 'varchar(50)', '操作模块'],
        ['operation', 'varchar(20)', '操作类型'],
        ['description', 'varchar(500)', '操作描述'],
        ['method', 'varchar(10)', '请求方法'],
        ['params', 'text', '请求参数'],
        ['ip', 'varchar(50)', 'IP地址'],
        ['oper_time', 'datetime', '操作时间'],
    ],
    '表3-7 操作日志表'
)

add_para('（8）字典表sys_dict，用于存储系统的字典数据，如表3-8所示：', first_indent=False)
add_table(
    ['字段名称', '数据类型', '字段说明'],
    [
        ['dict_id', 'bigint', '字典ID，主键，自动增长'],
        ['dict_type', 'varchar(50)', '字典类型，非空'],
        ['dict_label', 'varchar(100)', '字典标签，非空'],
        ['dict_value', 'varchar(100)', '字典值，非空'],
        ['sort', 'int', '排序'],
        ['status', 'tinyint', '状态（0正常 1停用），默认0'],
        ['remark', 'varchar(500)', '备注'],
    ],
    '表3-8 字典表'
)

doc.add_page_break()

# ========== 第4章 系统实现 ==========
add_heading('4.系统实现', 1)

add_heading('4.1系统主界面', 2)
add_para('用户成功登录系统后，进入系统首页（工作台界面）。首页采用卡片式布局，顶部展示门店的核心经营数据，包括寄养中宠物数量、待入店订单数量、本月订单数量和本月寄养金额四个统计卡片。中部左侧为今日待办订单列表，展示待入店和寄养中的订单信息，工作人员可以直接点击卡片跳转到订单详情页面进行操作；中部右侧为在店宠物状态概览，展示当前寄养中的宠物数量和状态信息。底部为系统快捷入口，提供新建寄养订单、新增客户、新增宠物档案、登记健康记录、登记服务记录、查看报表等常用操作的快捷按钮，点击快捷按钮可以直接打开对应功能的新增弹窗，提高操作效率。首页还配有温馨的宠物插画，营造温暖亲和的界面氛围。')
add_para('系统采用侧边栏加内容区域的布局结构，左侧侧边栏包含系统logo、导航菜单和用户信息区域。导航菜单按业务分类组织，分为工作台、业务管理（客户信息管理、宠物档案管理、寄养订单管理、健康档案管理、寄养服务记录）、数据统计（统计分析）和系统管理四个分组。侧边栏底部显示当前登录用户的头像、姓名和角色，点击用户区域可以展开操作菜单，进行角色切换和退出登录操作。系统的主色调采用暖橙色，搭配奶油白和浅杏色，营造温暖亲和的视觉风格，符合宠物寄养行业的品牌调性。')

add_heading('4.2客户信息管理模块', 2)
add_para('客户信息管理模块用于维护宠物主人的联系信息。页面顶部为搜索栏，支持按客户姓名和手机号码进行模糊搜索。搜索栏右侧为新增客户按钮，点击后弹出新增客户表单。页面主体为客户信息列表，采用表格形式展示客户的姓名、手机号码、地址、关联宠物数量、创建时间和操作列。列表支持分页显示，每页显示的数据条数根据屏幕高度动态计算，最后一页数据不足时自动补充空白行保持布局整齐。分页控件提供页码按钮、上一页和下一页按钮，当前页码高亮显示。')
add_para('新增和编辑客户使用同一个弹窗表单，表单包含姓名、手机号码、地址和备注四个字段。系统对手机号码进行格式校验和唯一性校验，如果手机号码格式不正确或已存在，表单将显示相应的错误提示并阻止提交。删除客户时，系统会先检查该客户名下是否有关联的宠物或订单，如果存在关联数据，系统将阻止删除操作并提示用户先处理关联数据，确保数据的完整性和一致性。客户信息的新增、编辑、删除操作都会被系统记录到操作日志中，便于后续的安全审计。')

add_heading('4.3宠物档案管理模块', 2)
add_para('宠物档案管理模块用于维护宠物的基本信息和健康档案。页面顶部为搜索栏，支持按宠物名称、品种和客户姓名进行搜索。页面主体为宠物信息列表，展示宠物的名称、物种、品种、性别、年龄、体重、所属客户、创建时间和操作列。宠物列表同样支持分页显示和动态每页条数。')
add_para('新增和编辑宠物档案使用弹窗表单，表单内容较为丰富，包括基本信息（宠物名称、物种、品种、性别、出生日期、体重、毛色）、健康信息（是否绝育、疫苗情况、过敏史）和特殊需求（特殊需求描述、备注）。宠物必须关联到一个已存在的客户，表单中提供客户选择下拉框。删除宠物时，系统检查该宠物是否有关联的寄养订单或健康记录，如果存在关联数据则阻止删除。宠物档案是健康记录的基础，每只宠物的健康记录都会关联到对应的宠物档案，形成完整的宠物健康数据积累。')

add_heading('4.4寄养订单管理模块', 2)
add_para('寄养订单管理模块是系统的核心模块，负责寄养订单的全生命周期管理。页面顶部为搜索栏，支持按订单编号、客户姓名、宠物名称和订单状态进行筛选。页面主体为订单列表，展示订单编号、客户姓名、宠物名称、寄养开始时间、寄养结束时间、总费用、订单状态和操作列。订单状态使用不同颜色的标签进行区分：待入店为橙色、寄养中为绿色、已完成为灰色、已取消为红色，便于工作人员快速识别订单状态。')
add_para('新增寄养订单使用弹窗表单，表单包含客户选择、宠物选择（根据所选客户动态加载该客户名下的宠物）、寄养开始时间、寄养结束时间、寄养类型、服务项目、总费用和备注。订单创建时，系统自动生成唯一的订单编号，编号格式为"FO+年月日+三位流水号"，如FO20261011001，便于订单的识别和追踪。')
add_para('订单的状态流转是本模块的核心功能。对于状态为"待入店"的订单，操作列提供"确认入店"按钮，点击后订单状态变更为"寄养中"，系统记录实际入店时间。对于状态为"寄养中"的订单，操作列提供"确认离店"按钮，点击后订单状态变更为"已完成"，系统记录实际离店时间。对于状态为"待入店"的订单，管理员还可以执行"取消订单"操作，订单状态变更为"已取消"，工作人员没有取消订单的权限。订单详情页面展示订单的完整信息，包括基本信息、关联的健康记录列表和服务记录列表，工作人员可以在订单详情页面直接登记健康记录和服务记录，系统自动关联到当前订单。')

add_heading('4.5健康档案管理模块', 2)
add_para('健康档案管理模块用于登记和查询宠物的健康记录。页面顶部为搜索栏，支持按宠物名称和记录时间范围进行筛选。页面主体为健康记录列表，展示记录时间、宠物名称、体温、体重、精神状态、食欲情况、异常描述、记录人和操作列。健康记录列表支持分页显示。')
add_para('新增健康记录使用弹窗表单，表单包含宠物选择、关联订单（可选，寄养期间的健康记录建议关联订单）、体温、体重、精神状态（下拉选择：良好/一般/较差）、食欲情况（下拉选择：正常/减少/不食）、排便情况（下拉选择：正常/异常）、异常描述、处理措施和记录时间。系统默认记录时间为当前时间，用户可以根据实际情况调整。为保证健康记录数据的准确性，系统限制只有当天的健康记录可以编辑，非当天的记录仅支持查看和删除，防止工作人员事后篡改历史记录。健康记录关联到宠物档案，在宠物的健康档案中可以查看该宠物的全部历史健康记录，形成完整的健康数据追踪。')

add_heading('4.6寄养服务记录模块', 2)
add_para('寄养服务记录模块用于登记寄养期间提供的附加服务信息。页面顶部为搜索栏，支持按订单编号和服务项目进行筛选。页面主体为服务记录列表，展示服务时间、订单编号、服务项目、服务内容、收费金额、操作人和操作列。')
add_para('新增服务记录使用弹窗表单，表单包含关联订单选择、服务项目、服务内容、收费金额和服务时间。服务记录必须关联到一个寄养订单，便于订单费用的核算。系统提供常用服务项目的字典管理，门店可以根据自身业务情况维护服务项目列表。服务记录的收费金额会在统计分析模块中纳入寄养收入的统计，为门店的经营分析提供数据支持。')

add_heading('4.7统计分析模块', 2)
add_para('统计分析模块为门店管理者提供多维度的经营数据可视化分析。页面顶部为筛选栏，提供周期筛选下拉（全部历史、近7天、近30天、本月、本季度、本年、自定义）和日期范围选择器，用户可以灵活选择统计的时间范围，默认显示从最早订单日期到今天的全部数据。')
add_para('页面主体包含四个统计卡片和两个图表。四个统计卡片分别展示订单总数、已完成订单、已取消订单和总营收，统计数据根据筛选的时间范围动态更新。两个图表分别为寄养订单趋势图和宠物品种分布环形图：寄养订单趋势图采用柱状图加折线图的组合形式，柱状图展示每月的订单数量，折线图展示每月的寄养收入，用户可以直观了解门店的订单量和收入趋势；宠物品种分布环形图展示在档宠物的品种占比，帮助门店了解客户宠物的品种构成。统计分析模块还支持订单数据的Excel导出功能，用户可以将订单数据导出为Excel文件，方便进行数据备份和进一步的离线分析。')

add_heading('4.8系统管理模块', 2)
add_para('系统管理模块用于管理系统的用户账号和权限，仅管理员角色可以访问。页面展示系统用户列表，包括账号、真实姓名、角色、账号状态、备注和操作列。管理员可以新增、编辑、删除用户账号，设置用户的角色（管理员或工作人员），启用或禁用用户账号。')
add_para('新增和编辑用户使用弹窗表单，包含账号、密码（编辑时可留空表示不修改）、真实姓名、角色、手机号码、邮箱和备注。系统对账号进行唯一性校验，密码在存储前使用BCrypt算法进行加密处理，确保密码安全。账号状态分为正常和禁用两种，被禁用的账号无法登录系统。管理员还可以为用户重置密码。系统管理模块的所有操作都会记录到操作日志中，包括操作人、操作时间、操作内容和IP地址，确保系统管理操作的可追溯性。此外，用户登录系统后可以在个人设置中修改自己的登录密码，修改密码时需要验证原密码，确保账号安全。')

doc.save(r'C:\Users\chen\Desktop\pet-foster\宠物寄养登记与健康档案管理系统_论文初稿.docx')
print('第二部分生成完成')
